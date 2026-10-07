"""Build GEB submission .docx with inline figures, Word (OMML) equations,
real tables and the reference list.
Source: GEB_main_blinded.md + GEB_figure_captions.md + bibliography.md;
figures inserted at first citation point in text."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
M = ROOT / "manuscript" / "GEB"
FIGDIR = ROOT / "figures"
OUT = M / "GEB_manuscript_inline_figures.docx"

FIGMAP = {
    "Figure 1": "fig_geb1_conceptual.png",
    "Figure 2": "fig_geb2_gamma_dominance.png",
    "Figure 3": "fig_geb3_2d_decay.png",
    "Figure 4": "fig_geb4_lumping_temporal.png",
    "Figure 5": "fig3_bias_heatmap.png",
    "Figure 6": "fig_geb6_empirical.png",
}
INSERT_AFTER = {
    "Figure 1": "Our objectives are",
    "Figure 2": "(Fig. 2a)",
    "Figure 3": "(Fig. 3)",
    "Figure 4": "Fig. 4b)",
    "Figure 5": "(Fig. 5)",
    "Figure 6": "(Fig. 6)",
}

# standalone equation lines -> OMML display equations
EQUATIONS = {
    "Δβ = β_P − β_H",
    "Bias = Δβ_obs − Δβ_true.",
    "Δβ_obs = Δβ_eco + B_γ + B_D + B_T + B_L + B_S + ε",
}

CAPTIONS = Path(M / "GEB_figure_captions.md").read_text()
cap = {}
for m in re.finditer(r"\*\*Figure (\d)\.\s*(.*?)\*\*\s*(.*?)(?=\n\*\*Figure|\Z)",
                     CAPTIONS, re.S):
    cap[f"Figure {m.group(1)}"] = (
        f"Figure {m.group(1)}. {m.group(2)} {m.group(3).strip()}")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)


def omml_run(text, italic=True):
    r = OxmlElement("m:r")
    if italic:
        stl = OxmlElement("m:sty")
        stl.set(qn("m:val"), "i")
        r.append(stl)
    t = OxmlElement("m:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def omml_ssub(base, sub):
    s = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(omml_run(base))
    sb = OxmlElement("m:sub")
    sb.append(omml_run(sub))
    s.append(e)
    s.append(sb)
    return s


def add_equation(text):
    """Render 'a = b_C − d_E' style equation as an OMML display equation."""
    p = doc.add_paragraph()
    p.alignment = 1
    omath = OxmlElement("m:oMathPara")
    om = OxmlElement("m:oMath")
    for tok in text.rstrip(".").split(" "):
        if "_" in tok:
            base, sub = tok.split("_", 1)
            om.append(omml_ssub(base, sub))
        elif tok in {"=", "+", "−", "-"}:
            om.append(omml_run(f" {tok} ", italic=False))
        else:
            om.append(omml_run(tok))
    omath.append(om)
    p._p.append(omath)


def add_fig(key):
    doc.add_picture(str(FIGDIR / FIGMAP[key]), width=Inches(6.0))
    doc.paragraphs[-1].alignment = 1
    p = doc.add_paragraph(cap.get(key, key))
    p.style = doc.styles["Intense Quote"] if "Intense Quote" in [
        s.name for s in doc.styles] else doc.styles["Normal"]
    p.runs[0].font.size = Pt(9)


def add_md_table(rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, celltxt in enumerate(row):
            c = tbl.cell(i, j)
            c.text = celltxt.strip()
            if i == 0:
                c.paragraphs[0].runs[0].bold = True
            for pr in c.paragraphs:
                for rn in pr.runs:
                    rn.font.size = Pt(9)


title = doc.add_heading(
    "Structural asymmetries distort guild-level beta-diversity contrasts", 0)
doc.add_paragraph("Blinded main text (double-anonymous).").italic = True

body = Path(M / "GEB_main_blinded.md").read_text()
lines = body.split("\n")
inserted = set()

# --- merge hard-wrapped lines into blocks: paragraphs, tables, equations ---
blocks = []
buf = []
tbl = []

def flush_buf():
    global buf
    if buf:
        blocks.append(("p", " ".join(buf)))
        buf = []

def flush_tbl():
    global tbl
    if tbl:
        rows = []
        for ln in tbl:
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if set("".join(cells)) <= {"-", ":"}:
                continue  # separator row
            rows.append(cells)
        if rows:
            blocks.append(("tbl", rows))
        tbl = []

for line in lines:
    s = line.strip()
    if s.startswith("|") and s.endswith("|"):
        flush_buf()
        tbl.append(s)
        continue
    flush_tbl()
    if not s:
        flush_buf()
        continue
    if s.startswith("#"):
        flush_buf()
        blocks.append(("h", s))
        continue
    if s in EQUATIONS:
        flush_buf()
        blocks.append(("eq", s))
        continue
    buf.append(s)
flush_buf()
flush_tbl()

for kind, s in blocks:
    if kind == "h":
        if s.startswith("# ") or s.startswith("*(blinded"):
            continue
        level = 1 if s.startswith("## ") else 2
        doc.add_heading(s.lstrip("#").strip(), level)
        continue
    if kind == "eq":
        add_equation(s)
        continue
    if kind == "tbl":
        add_md_table(s)
        continue
    doc.add_paragraph(s)
    for key, marker in INSERT_AFTER.items():
        if key not in inserted and marker in s:
            add_fig(key)
            inserted.add(key)

# --- references ---
doc.add_page_break()
doc.add_heading("References", 1)
bib = Path(M / "bibliography.md").read_text().splitlines()
for ln in bib:
    ln = ln.strip()
    if ln.startswith("- **"):
        m = re.match(r"- \*\*(.+?)\*\* — (.*)", ln)
        doc.add_paragraph(m.group(2) if m else ln[2:])

doc.save(str(OUT))
print("saved", OUT, "| inserted:", sorted(inserted))
